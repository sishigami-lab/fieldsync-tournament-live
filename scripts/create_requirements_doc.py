from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.section import WD_SECTION
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.style import WD_STYLE_TYPE
from pathlib import Path

OUT = Path(__file__).resolve().parents[1] / "deliverables" / "FieldSync_リアルタイム大会管理アプリ_要件定義書.docx"
OUT.parent.mkdir(exist_ok=True)

GREEN = "087F4F"; DARK = "173F31"; PALE = "EAF6EF"; INK = "17231D"; MUTED = "66756E"; LINE = "DCE5DF"; LIGHT = "F5F7F5"; RED = "A53A3A"
FONT = "Arial Unicode MS"

def set_font(run, size=10.5, bold=False, color=INK):
    run.font.name = FONT; run.font.size = Pt(size); run.bold = bold; run.font.color.rgb = RGBColor.from_string(color)
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), FONT)

def shade(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr(); shd = tcPr.find(qn("w:shd")) or OxmlElement("w:shd"); shd.set(qn("w:fill"), fill)
    if shd.getparent() is None: tcPr.append(shd)

def margins(cell, top=90, start=120, bottom=90, end=120):
    tc = cell._tc.get_or_add_tcPr(); tcMar = tc.first_child_found_in("w:tcMar")
    if tcMar is None: tcMar=OxmlElement("w:tcMar"); tc.append(tcMar)
    for side,val in (("top",top),("start",start),("bottom",bottom),("end",end)):
        node=tcMar.find(qn(f"w:{side}")) or OxmlElement(f"w:{side}"); node.set(qn("w:w"),str(val)); node.set(qn("w:type"),"dxa")
        if node.getparent() is None: tcMar.append(node)

def set_cell_text(cell, text, bold=False, color=INK, size=9.3, align=None):
    cell.text=""; p=cell.paragraphs[0]; p.paragraph_format.space_after=Pt(0); p.paragraph_format.line_spacing=1.15
    if align is not None: p.alignment=align
    set_font(p.add_run(str(text)), size, bold, color); cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER; margins(cell)

def set_table_widths(table, widths):
    table.autofit=False; total=sum(widths); table.alignment=WD_TABLE_ALIGNMENT.LEFT
    tblPr=table._tbl.tblPr; tblW=tblPr.first_child_found_in("w:tblW")
    if tblW is None: tblW=OxmlElement("w:tblW"); tblPr.append(tblW)
    tblW.set(qn("w:w"),str(total)); tblW.set(qn("w:type"),"dxa")
    ind=tblPr.first_child_found_in("w:tblInd")
    if ind is None: ind=OxmlElement("w:tblInd"); tblPr.append(ind)
    ind.set(qn("w:w"),"120"); ind.set(qn("w:type"),"dxa")
    grid=table._tbl.tblGrid
    for el in list(grid): grid.remove(el)
    for w in widths:
        col=OxmlElement("w:gridCol"); col.set(qn("w:w"),str(w)); grid.append(col)
    for row in table.rows:
        for i,(cell,w) in enumerate(zip(row.cells,widths)):
            cell.width=Inches(w/1440); tcW=cell._tc.get_or_add_tcPr().first_child_found_in("w:tcW")
            if tcW is None: tcW=OxmlElement("w:tcW"); cell._tc.get_or_add_tcPr().append(tcW)
            tcW.set(qn("w:w"),str(w)); tcW.set(qn("w:type"),"dxa")

def add_table(headers, rows, widths):
    t=doc.add_table(rows=1, cols=len(headers)); t.style="Table Grid"
    for i,h in enumerate(headers): set_cell_text(t.rows[0].cells[i],h,True,"FFFFFF",9,WD_ALIGN_PARAGRAPH.CENTER); shade(t.rows[0].cells[i],GREEN)
    for ri,row in enumerate(rows):
        cells=t.add_row().cells
        for i,v in enumerate(row): set_cell_text(cells[i],v, i==0, INK, 9); shade(cells[i],"FFFFFF" if ri%2==0 else LIGHT)
    set_table_widths(t,widths); doc.add_paragraph().paragraph_format.space_after=Pt(1)
    return t

def add_heading(text, level=1):
    p=doc.add_paragraph(text, style=f"Heading {level}"); p.paragraph_format.keep_with_next=True; return p

def add_p(text, bold_lead=None):
    p=doc.add_paragraph(); p.paragraph_format.space_after=Pt(6); p.paragraph_format.line_spacing=1.18
    if bold_lead and text.startswith(bold_lead):
        set_font(p.add_run(bold_lead),10.5,True); set_font(p.add_run(text[len(bold_lead):]),10.5)
    else: set_font(p.add_run(text),10.5)
    return p

def add_bullets(items):
    for item in items:
        p=doc.add_paragraph(style="Req Bullet"); set_font(p.add_run(item),10.2)

def callout(label, text, fill=PALE):
    t=doc.add_table(rows=1,cols=1); set_table_widths(t,[9360]); c=t.cell(0,0); shade(c,fill); c.text=""; p=c.paragraphs[0]; p.paragraph_format.space_after=Pt(0)
    set_font(p.add_run(label+"  "),10,True,GREEN if fill==PALE else RED); set_font(p.add_run(text),10,False,INK); margins(c,150,180,150,180); doc.add_paragraph().paragraph_format.space_after=Pt(1)

doc=Document(); sec=doc.sections[0]; sec.page_width=Inches(8.5); sec.page_height=Inches(11); sec.top_margin=Inches(.78); sec.bottom_margin=Inches(.75); sec.left_margin=Inches(1); sec.right_margin=Inches(1); sec.header_distance=Inches(.35); sec.footer_distance=Inches(.4)

styles=doc.styles
normal=styles["Normal"]; normal.font.name=FONT; normal.font.size=Pt(10.5); normal.font.color.rgb=RGBColor.from_string(INK); normal._element.rPr.rFonts.set(qn("w:eastAsia"),FONT); normal.paragraph_format.space_after=Pt(6); normal.paragraph_format.line_spacing=1.18
for name,size,before,after,color in [("Heading 1",17,15,8,GREEN),("Heading 2",13,11,5,DARK),("Heading 3",11,8,4,DARK)]:
    s=styles[name]; s.font.name=FONT; s.font.size=Pt(size); s.font.bold=True; s.font.color.rgb=RGBColor.from_string(color); s._element.rPr.rFonts.set(qn("w:eastAsia"),FONT); s.paragraph_format.space_before=Pt(before); s.paragraph_format.space_after=Pt(after); s.paragraph_format.keep_with_next=True
bullet=styles.add_style("Req Bullet",WD_STYLE_TYPE.PARAGRAPH); bullet.base_style=normal; bullet.paragraph_format.left_indent=Inches(.38); bullet.paragraph_format.first_line_indent=Inches(-.19); bullet.paragraph_format.space_after=Pt(4)
# real bullet numbering
numPart=doc.part.numbering_part.element; abstract=OxmlElement("w:abstractNum"); abstract.set(qn("w:abstractNumId"),"42"); lvl=OxmlElement("w:lvl"); lvl.set(qn("w:ilvl"),"0")
for tag,val in (("w:start","1"),("w:numFmt","bullet"),("w:lvlText","●"),("w:lvlJc","left")):
    e=OxmlElement(tag); e.set(qn("w:val"),val); lvl.append(e)
pPr=OxmlElement("w:pPr"); tabs=OxmlElement("w:tabs"); tab=OxmlElement("w:tab"); tab.set(qn("w:val"),"num"); tab.set(qn("w:pos"),"540"); tabs.append(tab); ind=OxmlElement("w:ind"); ind.set(qn("w:left"),"540"); ind.set(qn("w:hanging"),"270"); pPr.append(tabs); pPr.append(ind); lvl.append(pPr); abstract.append(lvl); numPart.append(abstract)
num=OxmlElement("w:num"); num.set(qn("w:numId"),"42"); aid=OxmlElement("w:abstractNumId"); aid.set(qn("w:val"),"42"); num.append(aid); numPart.append(num)
numPr=OxmlElement("w:numPr"); ilvl=OxmlElement("w:ilvl"); ilvl.set(qn("w:val"),"0"); nid=OxmlElement("w:numId"); nid.set(qn("w:val"),"42"); numPr.append(ilvl); numPr.append(nid); bullet._element.get_or_add_pPr().append(numPr)

# header/footer
hp=sec.header.paragraphs[0]; hp.text="FIELD SYNC  |  REQUIREMENTS DEFINITION"; hp.alignment=WD_ALIGN_PARAGRAPH.RIGHT; set_font(hp.runs[0],8,True,MUTED)
fp=sec.footer.paragraphs[0]; fp.alignment=WD_ALIGN_PARAGRAPH.CENTER; set_font(fp.add_run("FieldSync リアルタイム大会管理アプリ　要件定義書"),8,False,MUTED)

# cover
p=doc.add_paragraph(); p.paragraph_format.space_before=Pt(80); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; set_font(p.add_run("FIELD SYNC"),13,True,GREEN)
p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.space_before=Pt(12); p.paragraph_format.space_after=Pt(8); set_font(p.add_run("リアルタイム大会管理＆\n審判スコア入力アプリ"),28,True,DARK)
p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; set_font(p.add_run("要件定義書  — MVP / Version 1.0"),14,False,MUTED)
p=doc.add_paragraph(); p.paragraph_format.space_before=Pt(40); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; set_font(p.add_run("サッカー大会の作成・現場運営・保護者配信を\nスマートフォンひとつで完結するための統合Webアプリ"),11,False,INK)
t=doc.add_table(rows=4,cols=2); set_table_widths(t,[2300,7060]); labels=[("文書名","FieldSync リアルタイム大会管理アプリ 要件定義書"),("対象","Webアプリ MVP"),("作成日","2026年8月31日"),("ステータス","初版・要件確定用")]
for r,(a,b) in zip(t.rows,labels): set_cell_text(r.cells[0],a,True,GREEN,9); shade(r.cells[0],PALE); set_cell_text(r.cells[1],b,False,INK,9); shade(r.cells[1],"FFFFFF")
doc.add_page_break()

add_heading("1. 文書の目的",1)
add_p("本書は、サッカー大会の事前準備、当日の試合運営、審判によるスコア入力、トーナメント進行、保護者への大会案内配信を一つのWebアプリで行うための要件を定義する。MVPで画面上確認できる機能と、本番運用時に必要となる拡張要件を明確に分ける。")
callout("プロダクトの到達点", "管理者が大会を作成して概要と競技ルールを設定し、審判がスマホで試合結果を入力すると、保護者向け画面へ即時反映され、大会案内の配信とPDF作成まで連続して行えること。")
add_heading("2. 背景と解決する課題",1)
add_bullets(["紙・表計算・メッセージアプリに情報が分散し、転記と確認に時間がかかる。","審判結果を本部が再入力するため、トーナメント表や保護者への共有が遅れる。","屋外・移動中のスマホ操作では、細かな文字入力や小さなボタンが使いにくい。","大会案内文やPDFを毎回手作業で作成する負担が大きい。","管理機能を一般閲覧者から分離し、安全に運用する必要がある。"])
add_heading("3. 対象範囲",1)
add_table(["区分","対象機能","MVP状態"],[
    ("大会準備","大会作成、概要入力、競技ルール設定","画面実装済み"),("当日運営","試合選択、ストップウォッチ、得点・カード・PK入力","画面実装済み"),("進行管理","本戦・チャレンジカップ進出、スケジュール更新","ローカル連携実装済み"),("情報発信","保護者向け大会案内、文面コピー、印刷/PDF、公開操作","画面実装済み"),("権限・設定","管理者ログイン、パスワード変更、運用設定","画面実装済み"),("本番基盤","データベース、複数端末同期、ユーザー管理、監査ログ","将来拡張")],[1800,5360,2200])

add_heading("4. 利用者と権限",1)
add_table(["利用者","主な利用場面","権限"],[
    ("保護者・参加者","大会前／大会当日","大会案内、対戦日程、トーナメント表、結果の閲覧"),("審判","担当試合中","管理者ログイン後、担当試合のタイマー・得点・カード・PK入力、結果送信"),("大会管理者・本部","大会準備／当日運営","大会作成、概要・ルール設定、公開、PDF作成、パスワード・運用設定、全試合管理")],[1800,2600,4960])
add_heading("5. 全体業務フロー",1)
add_table(["段階","担当","操作","成果物・反映"],[
    ("1. 大会作成","管理者","新しい大会を作成","大会の下書き"),("2. 概要設定","管理者","日時、会場、注意事項等を入力","保護者向け案内の自動生成"),("3. 競技設定","管理者","試合時間、形式、PK人数を設定","審判画面へ即時反映"),("4. 確認・配信","管理者","プレビュー、PDF、公開・配信","閲覧用大会案内"),("5. 試合運営","審判","試合選択、計時、スコア入力、送信","スケジュールとトーナメント更新"),("6. 閲覧","保護者","公開画面を閲覧","最新の日程・結果を確認")],[1200,1400,3500,3260])

add_heading("6. 画面要件",1)
add_table(["画面ID","画面名","公開範囲","主な内容"],[
    ("SC-01","トーナメント表","全員","本戦／チャレンジ切替、試合カード、スコア、終了状態、進出先"),("SC-02","対戦スケジュール","全員","時間・コート別一覧、コート絞込、結果・終了表示"),("SC-03","大会案内","全員","日時、会場、当日の流れ、駐車場、雨天、医事、注意事項、表彰"),("SC-04","審判スコア入力","管理者・審判","担当試合、ストップウォッチ、得点、カード、PK、結果送信"),("SC-05","大会設定","管理者","大会作成、概要編集、競技ルール、確認、PDF、公開・配信"),("SC-06","システム設定","管理者","パスワード変更、確認・音・振動・通知の切替"),("SC-07","管理者ログイン","未ログイン者","数字4桁の簡易パスワード入力")],[1100,1900,1500,4860])

add_heading("7. 機能要件",1)
add_heading("7.1 管理者ログイン・パスワード",2)
add_table(["ID","要件","受入条件"],[
    ("AUTH-01","初期パスワードは1234とする。","正しい4桁入力で管理者画面が開く。"),("AUTH-02","審判・大会設定・システム設定は未ログイン時にロックする。","選択時にログイン画面が表示される。"),("AUTH-03","管理者は現在値を確認したうえで数字4桁へ変更できる。","現在値、新規値、確認値が正しい場合のみ変更される。"),("AUTH-04","変更後の値を次回ログイン判定に即時利用する。","ログアウト後、旧値では入れず新値で入れる。"),("AUTH-05","MVPでは状態を端末の画面表示中のみ保持する。","再読込で初期状態に戻ることを注記する。")],[1200,4600,3560])
add_heading("7.2 大会作成・概要設定",2)
add_table(["ID","要件","受入条件"],[
    ("EVT-01","管理者は新しい大会を作成できる。","新規作成後、概要入力状態へ遷移する。"),("EVT-02","大会名、開催日・時間、会場、住所、地図URL、主催、参加費、対象を設定できる。","入力値がプレビューへ反映される。"),("EVT-03","当日の流れ、駐車場、雨天対応、注意事項、医事・保険、表彰を設定できる。","複数行の内容が案内画面へ反映される。"),("EVT-04","作業導線は「大会作成→概要設定→確認・配信→公開完了」とする。","現在の工程が画面上で識別できる。")],[1200,4600,3560])
add_heading("7.3 競技レギュレーション",2)
add_table(["ID","要件","受入条件"],[
    ("RULE-01","試合時間は15・20・30分のクイック選択と、1〜120分の直接入力に対応する。","変更値が審判画面の基準時間へ反映される。"),("RULE-02","試合形式は前後半制／1本勝負を切り替えられる。","選択状態が画面に保持される。"),("RULE-03","PK規定人数は1・3・5人の選択と、1〜20人の直接入力に対応する。","規定人数終了後はサドンデス扱いとして表示する。")],[1200,4600,3560])
add_heading("7.4 審判スコア入力",2)
add_table(["ID","要件","受入条件"],[
    ("REF-01","未終了の対戦カードから担当試合をタップ選択できる。","選択した試合のチーム名・時刻・コートが表示される。"),("REF-02","時間表示は0:00から進むストップウォッチとする。","開始中は1秒ごとに加算し、停止中は変化しない。"),("REF-03","前半開始、ストップ、後半開始、試合終了を大きなボタンで操作できる。","操作に応じてフェーズ表示と計時状態が変わる。"),("REF-04","両チームの得点を大きな＋／−ボタンで入力できる。","0未満にならず、タップごとに即時表示される。"),("REF-05","イエロー／レッドカードをワンタップで加算できる。","チーム別の件数が即時表示される。"),("REF-06","PK戦モードを切り替え、オン時のみ成功数＋／−を表示する。","オフ時はPK入力欄を非表示にする。"),("REF-07","結果送信前にスコア確認を表示する。","勝者が確定していない場合は送信できない。")],[1200,4600,3560])
add_heading("7.5 自動連携",2)
add_table(["ID","要件","受入条件"],[
    ("SYNC-01","送信した試合を終了状態に更新する。","スケジュールにスコアと「終了」を表示する。"),("SYNC-02","1回戦勝者を本戦次ラウンドへ割り当てる。","該当する準決勝枠へチーム名が入る。"),("SYNC-03","1回戦敗者をチャレンジカップへ割り当てる。","該当する裏トーナメント枠へチーム名が入る。"),("SYNC-04","準決勝・チャレンジ準決勝の勝者を各決勝へ割り当てる。","該当する決勝枠へチーム名が入る。"),("SYNC-05","通常同点時はPK結果を用いて勝者を決定する。","PK差がない場合は勝者未確定とする。")],[1200,4600,3560])
add_heading("7.6 大会案内・配信・PDF",2)
add_table(["ID","要件","受入条件"],[
    ("PUB-01","設定内容から保護者向け大会案内画面を自動構成する。","入力内容と表示内容が一致する。"),("PUB-02","LINE・メール向け案内文を自動生成しコピーできる。","コピー完了が画面に表示される。"),("PUB-03","案内文をテキストとして保存できる。","大会名を含むファイル名で保存される。"),("PUB-04","印刷機能を利用してPDFを作成できる。","大会案内のみが読みやすく印刷される。"),("PUB-05","公開・配信前に最終確認を表示する。","確定後に公開中状態と完了通知が表示される。")],[1200,4600,3560])
add_heading("7.7 システム・運用設定",2)
add_table(["ID","設定項目","動作"],[
    ("SYS-01","結果送信前の確認","オン時は送信直前に確認画面を表示する。"),("SYS-02","操作音","オン時は得点・カード操作を音で通知する。本番版で実装する。"),("SYS-03","ボタン操作時の振動","オン時は対応端末で触覚フィードバックを行う。本番版で実装する。"),("SYS-04","ライブ更新のお知らせ","オン時は結果反映後に完了通知を表示する。")],[1200,3000,5160])

add_heading("8. UI／UX要件",1)
add_bullets(["スマートフォンの片手操作を優先し、審判画面では試合選択以外の文字入力を不要とする。","得点の＋／−、計時、結果送信は十分な高さと余白を確保し、隣接操作との誤タップを抑える。","屋外の太陽光下でも識別しやすい高コントラスト配色とする。","緑・白・ダークグレーを基調とし、成功・公開は緑、危険操作は赤で統一する。","PCでは上部ナビゲーション、スマホでは下部固定ナビゲーションを使用する。","管理者専用機能には鍵アイコンを表示し、権限の違いを視覚化する。","結果や公開処理の完了はトースト通知で即時にフィードバックする。"])
add_heading("9. データ要件",1)
add_table(["データ","主な項目","更新主体"],[
    ("大会","大会名、日時、会場、主催、対象、参加費、案内項目、公開状態","管理者"),("競技ルール","試合時間、試合形式、PK規定人数","管理者"),("試合","試合番号、時刻、コート、ラウンド、チーム、得点、PK、終了状態","審判／自動処理"),("カード","チーム別イエロー数・レッド数","審判"),("管理設定","管理者パスワード、確認・音・振動・通知設定","管理者")],[1800,5200,2360])
callout("MVPの制約", "現在の状態管理はReactローカルステートであり、画面再読込・別端末・別ブラウザには引き継がれない。本番運用ではサーバー保存とリアルタイム同期が必須となる。", "FFF3E6")

add_heading("10. 非機能要件",1)
add_table(["分類","MVP要件","本番化要件"],[
    ("対応端末","スマートフォン／タブレット／PCのレスポンシブ表示","主要iOS Safari・Android Chrome・PCブラウザで試験"),("性能","タップ操作を即時に画面反映","通常操作1秒以内、リアルタイム配信3秒以内を目標"),("可用性","単一ページで主要操作が継続可能","大会開催中の監視、バックアップ、障害時復旧"),("セキュリティ","4桁簡易パスワードによる画面ロック","個別アカウント、暗号化、権限、試行回数制限、監査ログ"),("データ保全","ローカル状態でデモ確認","DB永続化、自動保存、履歴、訂正・取消機能"),("アクセシビリティ","大きなタップ領域、色以外の状態表示","WCAG 2.2 AAを目安に検証"),("個人情報","個人名を扱わない大会・チーム情報中心","収集項目を最小化し、利用目的と保管期間を明示")],[1500,3520,4340])

add_heading("11. 技術要件",1)
add_table(["項目","採用内容"],[
    ("フレームワーク","Next.js（App Router）／TypeScript"),("UI","Tailwind CSS、lucide-react、レスポンシブデザイン"),("状態管理（MVP）","React Stateによるブラウザ内ローカル状態"),("画面構成","1ページ完結、タブ／ヘッダー切替"),("PDF","ブラウザ印刷機能を利用したPDF保存"),("本番候補","認証基盤、データベース、WebSocketまたは購読型リアルタイム更新、ファイル保存")],[2100,7260])

add_heading("12. 受入テスト",1)
add_table(["No.","テストシナリオ","期待結果"],[
    ("AT-01","未ログインで審判画面を選択","パスワード入力画面が表示される。"),("AT-02","1234で初回ログイン","管理者モードになり大会設定が開く。"),("AT-03","パスワードを4321へ変更してログアウト","旧番号では入れず、新番号で再ログインできる。"),("AT-04","試合時間を任意値へ変更","設定値が保持され審判画面の表示へ反映される。"),("AT-05","ストップウォッチを開始・停止・再開","経過時間が正しく加算／停止される。"),("AT-06","得点・カード・PKを入力","各値が0未満にならず、タップごとに反映される。"),("AT-07","1回戦結果を送信","スケジュールが終了表示になり、勝者・敗者が所定枠へ進む。"),("AT-08","大会概要を編集して案内を開く","入力内容から案内画面と送信用文章が生成される。"),("AT-09","PDF作成を選択","印刷画面から大会案内をPDF保存できる。"),("AT-10","スマホ幅で全画面を操作","下部メニューと主要ボタンが重ならず操作できる。")],[1000,4300,4060])

add_heading("13. 本番化に向けた優先拡張",1)
add_table(["優先度","拡張項目","理由"],[
    ("P0","データベース永続化・複数端末リアルタイム同期","現場の複数審判・本部・保護者で同じ結果を共有するため。"),("P0","管理者・審判アカウントと役割別権限","4桁共通番号より安全に担当範囲を制御するため。"),("P0","結果訂正・取消・操作履歴","誤入力対応と大会結果の信頼性を確保するため。"),("P1","チーム・対戦カード・コート・時刻の管理画面","大会ごとに組合せと日程を自由作成するため。"),("P1","通知配信連携（メール／LINE等）","公開後の案内と変更を確実に届けるため。"),("P1","サーバー側PDF生成と保管","端末差なく同じ書式で配布できるようにするため。"),("P2","オフライン入力・復帰時同期","通信が不安定なグラウンドで入力を継続するため。"),("P2","大会複製・テンプレート","定期大会の準備時間を短縮するため。")],[1100,3600,4660])

add_heading("14. 前提・決定事項",1)
add_bullets(["MVPは実際の操作感と業務フローの確認を目的とし、React Stateで動作する。","同点時の勝者決定にはPK結果を使用し、通常得点・PKとも同点の場合は送信不可とする。","チャレンジカップは1回戦敗退チームを自動割当する。","管理者パスワードは数字4桁とし、初期値は1234とする。","大会案内のPDFはMVPではブラウザ印刷、本番ではサーバー生成を検討する。","公開範囲、個人情報、運用責任者、バックアップ方針は本番開発開始前に確定する。"])
add_heading("15. 完了条件",1)
add_bullets(["本書のAT-01〜AT-10を主要対象端末で実施し、重大な不具合がない。","大会作成から案内公開、試合結果送信、トーナメント反映まで一連のデモが完走できる。","管理者専用画面が未ログイン利用者から保護される。","スマホで得点・カード・計時・送信を片手で操作できる。","MVP制約と本番化対象について関係者の合意が得られる。"])

doc.core_properties.title="FieldSync リアルタイム大会管理アプリ 要件定義書"
doc.core_properties.subject="サッカー大会管理・審判スコア入力アプリ MVP"
doc.core_properties.author="FieldSync Project"
doc.save(OUT)
print(OUT)
